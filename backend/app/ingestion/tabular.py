"""Spreadsheet and delimited-file parsing.

An honest caveat before the code: **semantic search over tabular data is
weak**, and no amount of chunking fixes that. "Which region had the highest
Q3 revenue?" is an aggregation, not a similarity lookup — the answer exists in
no single row, so no retrieved row contains it. The right tool for that
question is SQL over the table, not embeddings over its rows.

What this does support well is *lookup*: "what was the Q3 revenue for the
North region?", "which product has SKU 4471?". Those questions match a
specific row, and row-per-chunk retrieval answers them accurately.

Supporting these formats is still worth it — spreadsheets are where
organisations keep their reference data — but the limitation is real and the
UI should not pretend otherwise.
"""

import csv
import io

import openpyxl

from app.core.logging import get_logger
from app.ingestion.archives import guard_ooxml
from app.ingestion.base import DocumentParser, ParsedDocument, ParsedPage, UnparsableDocumentError
from app.ingestion.tables import MAX_ROWS, format_value, looks_like_header, render_row

logger = get_logger(__name__)

_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


class XlsxParser(DocumentParser):
    content_types = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",)
    extensions = (".xlsx", ".xlsm")

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        guard_ooxml(data, filename=filename)

        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(data),
                # Streams rows instead of building the whole sheet in memory —
                # essential for large exports.
                read_only=True,
                # Cached values rather than formula strings. Without this, a
                # cell reads as "=SUM(B2:B99)", which is useless to embed and
                # useless to answer from.
                data_only=True,
            )
        except Exception as exc:  # noqa: BLE001
            raise UnparsableDocumentError(
                f"Could not read {filename} as a spreadsheet. "
                "If it is an older .xls file, save it as .xlsx first."
            ) from exc

        pages: list[ParsedPage] = []
        try:
            # One sheet per page: a natural, meaningful citation target, and
            # it keeps unrelated tables from being chunked together.
            for number, worksheet in enumerate(workbook.worksheets, start=1):
                text = _render_sheet(worksheet)
                if text:
                    pages.append(
                        ParsedPage(
                            number=number,
                            text=f"# Sheet: {worksheet.title}\n\n{text}",
                            metadata={"sheet_name": worksheet.title},
                        )
                    )
        finally:
            # read_only mode holds the archive open until closed explicitly.
            workbook.close()

        if not pages:
            raise UnparsableDocumentError(f"{filename} contains no readable data.")

        logger.info("xlsx_parsed", extra={"source_name": filename, "sheets": len(pages)})

        return ParsedDocument(
            pages=pages,
            metadata={"source_page_count": len(pages)},
        )


def _render_sheet(worksheet) -> str:  # noqa: ANN001 - openpyxl worksheet type is awkward
    rows = worksheet.iter_rows(values_only=True)

    try:
        first = next(rows)
    except StopIteration:
        return ""

    first_values = list(first)
    if looks_like_header(first_values):
        headers = [format_value(value) for value in first_values]
        lines: list[str] = []
    else:
        headers = []
        lines = [line] if (line := render_row([], first_values)) else []

    count = 0
    for values in rows:
        if count >= MAX_ROWS:
            lines.append("[further rows omitted]")
            break
        rendered = render_row(headers, list(values))
        if rendered:
            lines.append(rendered)
            count += 1

    return "\n".join(lines)


class CsvParser(DocumentParser):
    content_types = ("text/csv", "application/csv", "text/tab-separated-values")
    extensions = (".csv", ".tsv")

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        text = _decode(data, filename)
        if not text.strip():
            raise UnparsableDocumentError(f"{filename} is empty.")

        # Sniff the delimiter rather than assuming a comma: European exports
        # commonly use semicolons, and .tsv uses tabs.
        sample = text[:8192]
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
                sample, delimiters=",;\t|"
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(io.StringIO(text), dialect)

        try:
            first = next(reader)
        except StopIteration:
            raise UnparsableDocumentError(f"{filename} is empty.") from None

        if looks_like_header(list(first)):
            headers = [format_value(value) for value in first]
            lines: list[str] = []
        else:
            headers = []
            lines = [line] if (line := render_row([], list(first))) else []

        truncated = False
        for index, values in enumerate(reader):
            if index >= MAX_ROWS:
                truncated = True
                break
            rendered = render_row(headers, list(values))
            if rendered:
                lines.append(rendered)

        if truncated:
            lines.append("[further rows omitted]")

        if not lines:
            raise UnparsableDocumentError(f"No rows could be read from {filename}.")

        logger.info("csv_parsed", extra={"source_name": filename, "rows": len(lines)})

        return ParsedDocument(
            pages=[ParsedPage(number=1, text="\n".join(lines))],
            metadata={"source_page_count": 1, "columns": headers or None},
        )


def _decode(data: bytes, filename: str) -> str:
    for encoding in _ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnparsableDocumentError(f"Could not decode {filename} as text.")
